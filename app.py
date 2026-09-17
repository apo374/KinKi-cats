from flask import Flask, jsonify, request, render_template, session, redirect, url_for
from urllib.parse import urlparse, urljoin, quote
from functools import wraps
from werkzeug.utils import secure_filename
import json
import math
import os
import time
import urllib.request
from datetime import datetime, timedelta, timezone

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    load_workbook = None

# app.py はプロジェクト直下に置く。
# 実体（templates / static / data）は bousai_app/ 配下にあるので、そこを参照する。
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
APP_DIR = os.path.join(BASE_DIR, 'bousai_app')

app = Flask(
    __name__,
    template_folder=os.path.join(APP_DIR, 'templates'),
    static_folder=os.path.join(APP_DIR, 'static'),
)
app.secret_key = 'your-secret-key-here'

# 管理者認証情報
ADMIN_CREDENTIALS = {
    'admin': '123'
}

# ────────────────────────────────
# 気象警報・注意報設定
PREFECTURE_CODE = "020000"  # 青森県
AREA_NAME = "青森市"

# 気象庁の警報・注意報JSONにおける青森市の市区町村コード
AREA_CODE = "0220100"

WARNING_URL = (
    f"https://www.jma.go.jp/bosai/warning/data/r8/{PREFECTURE_CODE}.json"
)

JST = timezone(timedelta(hours=9))
WEATHER_CACHE_SECONDS = 600
_weather_cache = None
_weather_cache_at = 0.0

# 警報・注意報のコード一覧
WARNING_CODES = {
    "00": "解除",
    "02": "暴風雪警報",
    "03": "レベル3大雨警報",
    "04": "洪水警報",
    "05": "暴風警報",
    "06": "大雪警報",
    "07": "波浪警報",
    "08": "レベル3高潮警報",
    "09": "レベル3土砂災害警報",
    "10": "レベル2大雨注意報",
    "12": "大雪注意報",
    "13": "風雪注意報",
    "14": "雷注意報",
    "15": "強風注意報",
    "16": "波浪注意報",
    "17": "融雪注意報",
    "18": "洪水注意報",
    "19": "レベル2高潮注意報",
    "20": "濃霧注意報",
    "21": "乾燥注意報",
    "22": "なだれ注意報",
    "23": "低温注意報",
    "24": "霜注意報",
    "25": "着氷注意報",
    "26": "着雪注意報",
    "27": "その他の注意報",
    "29": "レベル2土砂災害注意報",
    "32": "暴風雪特別警報",
    "33": "レベル5大雨特別警報",
    "35": "暴風特別警報",
    "36": "大雪特別警報",
    "37": "波浪特別警報",
    "38": "レベル5高潮特別警報",
    "39": "レベル5土砂災害特別警報",
    "43": "レベル4大雨危険警報",
    "48": "レベル4高潮危険警報",
    "49": "レベル4土砂災害危険警報"
}

# ────────────────────────────────
# サンプルデータの読み込み
DATA_FILE = os.path.join(APP_DIR, 'data', 'shelters.json')
INSTRUCTIONS_FILE = os.path.join(APP_DIR, 'data', 'instructions.json')
BOARD_DATA_FILE = os.path.join(APP_DIR, 'data', 'board_data.xlsx')
BOARD_SHEETS = {
    '被害状況地図': ['地図'],
    '被害状況一覧': ['場所', '被害内容', '発生日時', '状況'],
    '避難指示': ['対象地域', '指示内容', '発令日時', '状況'],
    '派遣職員': ['氏名', '所属', '派遣先', '派遣日時'],
}
BOARD_LOCATION_COORDS = {
    '青森市中央部': [40.8222, 140.7474],
    '青森駅周辺': [40.8298, 140.7346],
    '浪岡地区': [40.7105, 140.5908],
    '浅虫地区': [40.8898, 140.8615],
    '八甲田周辺': [40.6792, 140.9318],
    '油川地区': [40.8580, 140.6860],
}
MAP_CENTER = [40.8244, 140.7400]

def load_json(path, default):
    """JSONファイルを読み込む（存在しない・壊れている場合は default を返す）"""
    try:
        with open(path, encoding='utf-8') as f:
            return json.load(f)
    except (FileNotFoundError, json.JSONDecodeError):
        return default

shelters = load_json(DATA_FILE, [])
instructions = load_json(INSTRUCTIONS_FILE, [])
_geocode_cache = {}

def save_instructions():
    """指示ボードのデータをファイルに保存する"""
    try:
        with open(INSTRUCTIONS_FILE, 'w', encoding='utf-8') as f:
            json.dump(instructions, f, ensure_ascii=False, indent=2)
    except Exception:
        pass


def save_shelters():
    """避難所データをファイルに保存する"""
    with open(DATA_FILE, 'w', encoding='utf-8') as f:
        json.dump(shelters, f, ensure_ascii=False, indent=2)


def _normalize_board_header(value):
    label = str(value or '').strip()
    if label == '地点ごとの被害状況':
        return '地図'
    return label


def _stringify_board_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y年%m月%d日 %H:%M')
    return str(value).strip()


def load_board_data():
    """Excel の指示・発信ボードデータを読み込む。失敗時は空データを返す。"""
    empty_data = {sheet_name: [] for sheet_name in BOARD_SHEETS}
    if load_workbook is None or not os.path.exists(BOARD_DATA_FILE):
        return empty_data

    try:
        workbook = load_workbook(BOARD_DATA_FILE, read_only=True, data_only=True)
    except Exception:
        return empty_data

    data = {}
    for sheet_name, required_headers in BOARD_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            data[sheet_name] = []
            continue

        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            data[sheet_name] = []
            continue

        headers = [_normalize_board_header(cell) for cell in rows[0]]
        lookup = {header: idx for idx, header in enumerate(headers) if header}
        records = []
        for row in rows[1:]:
            if row is None or not any(cell is not None and str(cell).strip() for cell in row):
                continue
            record = {}
            for required in required_headers:
                if required not in lookup:
                    continue
                record[required] = _stringify_board_value(row[lookup[required]])
            if record and any(str(value).strip() for value in record.values()):
                records.append(record)
        data[sheet_name] = records

    workbook.close()
    return data


def build_damage_map_points(rows):
    """被害状況一覧から地図マーカー用のデータを作る。対応表にない場所は除外する。"""
    points = []
    for row in rows or []:
        place = _stringify_board_value(row.get('場所', '')).strip()
        if not place or place not in BOARD_LOCATION_COORDS:
            continue
        coords = BOARD_LOCATION_COORDS[place]
        points.append({
            'location': place,
            'content': _stringify_board_value(row.get('被害内容', '')),
            'status': _stringify_board_value(row.get('状況', '')),
            'datetime': _stringify_board_value(row.get('発生日時', '')),
            'latitude': coords[0],
            'longitude': coords[1],
        })
    return points


SHELTER_UPLOAD_DIR = os.path.join(APP_DIR, 'static', 'uploads', 'shelters')
os.makedirs(SHELTER_UPLOAD_DIR, exist_ok=True)
ALLOWED_IMAGE_MIME_TYPES = {'image/jpeg', 'image/png', 'image/gif', 'image/webp'}
MAX_UPLOAD_BYTES = 5 * 1024 * 1024


def get_shelter_by_id(shelter_id):
    """IDで避難所を検索する"""
    try:
        target_id = int(shelter_id)
    except (TypeError, ValueError):
        return None
    for shelter in shelters:
        if isinstance(shelter, dict) and shelter.get('id') == target_id:
            return shelter
    return None


def is_local_upload_path(image_path):
    """アップロード済みイメージがアプリ内のローカル保存先か確認する"""
    if not image_path or not isinstance(image_path, str):
        return False
    normalized = image_path.strip().replace('\\', '/')
    if normalized.startswith('http://') or normalized.startswith('https://'):
        return False
    return '/uploads/' in normalized or normalized.startswith('uploads/')


def remove_uploaded_image(image_path):
    """管理対象画像を削除する"""
    if not is_local_upload_path(image_path):
        return
    relative = image_path.lstrip('/')
    if relative.startswith('uploads/'):
        file_path = os.path.join(APP_DIR, 'static', relative)
    else:
        file_path = os.path.join(APP_DIR, 'static', relative.replace('uploads/', ''))
    try:
        if os.path.isfile(file_path):
            os.remove(file_path)
    except OSError:
        pass


def save_uploaded_image(uploaded_file):
    """画像を安全に保存して、保存先の相対URLを返す"""
    if not uploaded_file or not uploaded_file.filename:
        return ''

    filename = secure_filename(uploaded_file.filename)
    if not filename:
        return ''

    mime_type = uploaded_file.mimetype or ''
    if mime_type not in ALLOWED_IMAGE_MIME_TYPES:
        raise ValueError('画像形式が不正です。JPEG、PNG、GIF、WebPのみ対応しています。')

    if uploaded_file.content_length and uploaded_file.content_length > MAX_UPLOAD_BYTES:
        raise ValueError('画像サイズが大きすぎます。5MB以内にしてください。')

    file_stream = uploaded_file.stream.read()
    if len(file_stream) > MAX_UPLOAD_BYTES:
        raise ValueError('画像サイズが大きすぎます。5MB以内にしてください。')

    _, ext = os.path.splitext(filename)
    if not ext:
        ext = '.png'
    safe_name = f"{int(time.time() * 1000)}_{secure_filename(os.path.splitext(filename)[0])}{ext.lower()}"
    save_path = os.path.join(SHELTER_UPLOAD_DIR, safe_name)
    with open(save_path, 'wb') as f:
        f.write(file_stream)
    return f"/uploads/shelters/{safe_name}"


def normalize_yes_no(value):
    if value is None:
        return 'なし'
    return 'あり' if str(value).lower() in {'1', 'true', 'yes', 'on', 'あり'} else 'なし'


# ────────────────────────────────

# ────────────────────────────────
# 認証関連の設定とヘルパー関数
def is_safe_url(target):
    """リダイレクト先URLが安全かどうかチェック"""
    ref_url = urlparse(request.host_url)
    test_url = urlparse(urljoin(request.host_url, target))
    return test_url.scheme in ('http', 'https') and ref_url.netloc == test_url.netloc

def login_required(f):
    """認証が必要なページに付けるデコレータ"""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get('logged_in'):
            # 現在のURLをnextパラメータとしてログイン画面にリダイレクト
            return redirect(url_for('login', next=request.url))
        return f(*args, **kwargs)
    return decorated_function

def get_japan_time():
    """日本時間（JST）の現在時刻を取得する"""
    return datetime.now(JST).strftime("%Y年%m月%d日 %H:%M")


def format_report_time(iso_str):
    """気象庁の発表時刻（ISO形式）をJSTの表示用文字列に変換する"""
    if not iso_str:
        return "不明"
    try:
        parsed = datetime.fromisoformat(iso_str.replace('Z', '+00:00'))
        if parsed.tzinfo:
            parsed = parsed.astimezone(JST)
        return parsed.strftime("%Y年%m月%d日 %H:%M")
    except ValueError:
        return iso_str


def filter_shelters(district=None):
    """district 指定があれば一致する避難所のみ、なければ全件を返す"""
    return [s for s in shelters if not district or s.get('district') == district]


def search_shelters(filters=None):
    """検索条件に合う避難所だけを返す。公開検索用の軽量フィルタ。"""
    if filters is None:
        filters = {}

    name = str(filters.get('name', '') or '').strip()
    postal_code = str(filters.get('postal_code', '') or '').strip()
    address = str(filters.get('address', '') or '').strip()
    capacity = str(filters.get('capacity', '') or '').strip()
    pet_allowed = filters.get('pet_allowed') or filters.get('pet-allowed')
    barrier_free = filters.get('barrier_free') or filters.get('barrier-free')

    def to_int(value):
        try:
            return int(value)
        except (TypeError, ValueError):
            return None

    min_capacity = to_int(capacity)
    normalized_postal = ''.join(ch for ch in postal_code if ch.isdigit())

    results = []
    for shelter in shelters:
        if not isinstance(shelter, dict):
            continue

        shelter_name = str(shelter.get('name', '') or '')
        shelter_postal = str(shelter.get('postal_code', '') or '')
        shelter_address = str(shelter.get('address', '') or '')
        shelter_capacity = to_int(shelter.get('capacity'))

        if name and name not in shelter_name:
            continue
        if normalized_postal:
            shelter_digits = ''.join(ch for ch in shelter_postal if ch.isdigit())
            if normalized_postal not in shelter_digits:
                continue
        if address:
            combined = f"{shelter_address} {shelter_postal}".strip()
            if address not in combined:
                continue
        if min_capacity is not None:
            if shelter_capacity is None or shelter_capacity < min_capacity:
                continue
        if pet_allowed == 'yes' and shelter.get('pet_allowed') is not True:
            continue
        if pet_allowed == 'no' and shelter.get('pet_allowed') is True:
            continue
        if barrier_free == 'yes':
            if shelter.get('barrier_free') not in (True, 'あり', '対応している'):
                continue
        if barrier_free == 'no':
            if shelter.get('barrier_free') in (True, 'あり', '対応している'):
                continue

        results.append(shelter)
    return results


def is_valid_coordinate(latitude, longitude):
    """緯度・経度が地図に描画できる数値か確認する"""
    try:
        latitude = float(latitude)
        longitude = float(longitude)
    except (TypeError, ValueError):
        return False
    return (math.isfinite(latitude) and math.isfinite(longitude)
            and -90 <= latitude <= 90 and -180 <= longitude <= 180)


def get_map_data():
    """地図に描画可能な避難所と被害・発信情報を返す"""
    current_shelters = load_json(DATA_FILE, [])
    current_instructions = load_json(INSTRUCTIONS_FILE, [])
    map_shelters = [
        {
            'id': shelter.get('id'),
            'name': shelter.get('name', '名称未設定'),
            'place_name': shelter.get('place_name', ''),
            'latitude': float(shelter['latitude']),
            'longitude': float(shelter['longitude']),
            'status': shelter.get('status', '利用状況不明')
        }
        for shelter in current_shelters
        if isinstance(shelter, dict)
        if is_valid_coordinate(shelter.get('latitude'), shelter.get('longitude'))
    ]
    incidents = [
        {
            'id': notice.get('id'),
            'title': notice.get('content', '発信情報'),
            'shelter': notice.get('shelter', ''),
            'status': notice.get('status', ''),
            'place_name': notice.get('place_name', ''),
            'updated_at': notice.get('updated_at', ''),
            'latitude': float(notice['latitude']),
            'longitude': float(notice['longitude'])
        }
        for notice in current_instructions
        if isinstance(notice, dict) and notice.get('target') == '住民'
        if is_valid_coordinate(notice.get('latitude'), notice.get('longitude'))
        and notice.get('status') not in ('解除', '完了')
    ]
    return {'shelters': map_shelters, 'incidents': incidents}


@app.route('/api/map_data')
def api_map_data():
    response = jsonify(get_map_data())
    response.headers['Cache-Control'] = 'no-store'
    return response


@app.route('/api/geocode')
def api_geocode():
    """住所検索を気象庁以外の外部サービスへ中継する"""
    query = request.args.get('q', '').strip()
    if not query:
        return jsonify({'error': '検索語を入力してください。'}), 400
    if len(query) > 100:
        return jsonify({'error': '検索語が長すぎます。'}), 400

    if query in _geocode_cache:
        return jsonify(_geocode_cache[query])

    url = (
        'https://nominatim.openstreetmap.org/search?format=jsonv2&limit=5&accept-language=ja&q='
        + quote(query)
    )
    try:
        request_obj = urllib.request.Request(
            url,
            headers={'User-Agent': 'bousai-app/1.0 contact: local-development'}
        )
        with urllib.request.urlopen(request_obj, timeout=8) as response:
            results = json.loads(response.read())
        locations = [
            {
                'name': result.get('display_name', query),
                'latitude': float(result['lat']),
                'longitude': float(result['lon'])
            }
            for result in results
            if is_valid_coordinate(result.get('lat'), result.get('lon'))
        ]
        payload = {'results': locations}
        _geocode_cache[query] = payload
        return jsonify(payload)
    except (OSError, ValueError, json.JSONDecodeError):
        return jsonify({'error': '場所を検索できませんでした。'}), 502


@app.route('/api/address_search')
def api_address_search():
    """郵便番号から住所を解決する公開API。検索画面から利用する。"""
    postal_code = (request.args.get('postal_code') or '').strip()
    if not postal_code:
        return jsonify({'error': '郵便番号を入力してください。'}), 400

    normalized = ''.join(ch for ch in postal_code if ch.isdigit())
    if len(normalized) != 7:
        return jsonify({'error': '郵便番号を7桁で入力してください。'}), 400

    url = f'https://zipcloud.ibsnet.co.jp/api/search?zipcode={normalized}'
    try:
        request_obj = urllib.request.Request(url, headers={'User-Agent': 'bousai-app/1.0'})
        with urllib.request.urlopen(request_obj, timeout=8) as response:
            payload = json.loads(response.read())
    except (OSError, ValueError, json.JSONDecodeError):
        return jsonify({'error': '住所検索に失敗しました。'}), 502

    results = payload.get('results')
    if not results:
        return jsonify({'error': '住所が見つかりませんでした。'}), 404

    first = results[0]
    address_parts = [
        first.get('address1', ''),
        first.get('address2', ''),
        first.get('address3', ''),
        first.get('address4', ''),
    ]
    full_address = ' '.join(part for part in address_parts if part).strip().replace('  ', ' ')
    if not full_address:
        return jsonify({'error': '住所が見つかりませんでした。'}), 404

    return jsonify({'address': full_address})


@app.route('/api/postal_code_lookup')
@login_required
def api_postal_code_lookup():
    """郵便番号から住所を取得するサーバー側プロキシ"""
    zipcode = request.args.get('zipcode', '').strip()
    if not zipcode or not zipcode.isdigit() or len(zipcode) != 7:
        return jsonify({'error': '郵便番号を7桁で入力してください。'}), 400

    url = f'https://zipcloud.ibsnet.co.jp/api/search?zipcode={zipcode}'
    try:
        request_obj = urllib.request.Request(url, headers={'User-Agent': 'bousai-app/1.0'})
        with urllib.request.urlopen(request_obj, timeout=8) as response:
            payload = json.loads(response.read())
    except (OSError, ValueError, json.JSONDecodeError):
        return jsonify({'error': '郵便番号の取得に失敗しました。通信状態を確認してください。'}), 502

    result = payload.get('results')
    if not result or not isinstance(result, list):
        return jsonify({'error': '該当する住所は見つかりませんでした。'}), 404

    first = result[0]
    address_parts = [
        first.get('address1', ''),
        first.get('address2', ''),
        first.get('address3', ''),
        first.get('address4', ''),
    ]
    full_address = ' '.join(part for part in address_parts if part).strip().replace('  ', ' ')
    if not full_address:
        return jsonify({'error': '該当する住所は見つかりませんでした。'}), 404

    return jsonify({'address': full_address})


def parse_area_warnings(warning_data):
    """気象庁JSONの最新報から対象市区町村の発表・継続中の情報を抽出する"""
    if not isinstance(warning_data, list):
        raise ValueError("気象庁の警報・注意報データが新形式の配列ではありません")

    def report_datetime(report):
        value = report.get("reportDatetime", "")
        try:
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        except (AttributeError, ValueError):
            return datetime.min.replace(tzinfo=timezone.utc)

    relevant_reports = []
    for report in warning_data:
        if not isinstance(report, dict):
            continue
        warning = report.get("warning")
        if not isinstance(warning, dict):
            continue
        class20_items = warning.get("class20Items", [])
        if not isinstance(class20_items, list):
            continue
        if any(
            isinstance(item, dict) and item.get("areaCode") == AREA_CODE
            for item in class20_items
        ):
            relevant_reports.append(report)

    if not relevant_reports:
        return [], ""

    latest_datetime = max(report_datetime(report) for report in relevant_reports)
    latest_reports = [
        report for report in relevant_reports
        if report_datetime(report) == latest_datetime
    ]
    warnings = []
    seen_codes = set()

    for report in latest_reports:
        area = next(
            item for item in report["warning"]["class20Items"]
            if isinstance(item, dict) and item.get("areaCode") == AREA_CODE
        )
        for kind in area.get("kinds", []):
            if not isinstance(kind, dict):
                continue
            status = kind.get("status", "").strip()
            code = kind.get("code", "")
            if status not in ("発表", "継続") or not code or code in seen_codes:
                continue
            warnings.append({
                "name": WARNING_CODES.get(
                    code,
                    f"不明な警報・注意報 (コード: {code})"
                ),
                "code": code,
                "status": status
            })
            seen_codes.add(code)

    return warnings, latest_datetime.isoformat()


def get_weather_warnings(force_refresh=False):
    """対象市区町村の警報・注意報を取得する"""
    global _weather_cache, _weather_cache_at

    now = time.monotonic()
    if not force_refresh and _weather_cache is not None and now - _weather_cache_at < WEATHER_CACHE_SECONDS:
        return _weather_cache

    try:
        # 青森県の新形式（令和8年～）警報・注意報データを取得
        with urllib.request.urlopen(url=WARNING_URL, timeout=10) as res:
            warning_data = json.loads(res.read())

        warnings, report_datetime = parse_area_warnings(warning_data)

        result = {
            "area_name": AREA_NAME,
            "warnings": warnings,
            "report_time": format_report_time(report_datetime),
            "last_fetch_time": get_japan_time()
        }
        _weather_cache = result
        _weather_cache_at = now
        return result

    except Exception:
        return {
            "area_name": AREA_NAME,
            "warnings": [],
            "report_time": "取得失敗",
            "last_fetch_time": get_japan_time(),
            "error": True
        }


# トップページ：templates/index.html を返す（住民向け指示も表示する）
@app.route('/')
def index():
    resident_notices = [
        i for i in load_json(INSTRUCTIONS_FILE, []) if i.get('target') == '住民'
    ]
    return render_template('index.html', resident_notices=resident_notices)

# ログインページ
@app.route('/login', methods=['GET', 'POST'])
def login():
    # リダイレクト先を取得（デフォルトは避難所登録画面）
    next_url = request.args.get('next') or request.form.get('next')

    # 安全でないURLの場合はデフォルトページにリダイレクト
    if not next_url or not is_safe_url(next_url):
        next_url = url_for('shelter_register')

    if request.method == 'POST':
        password = request.form.get('password', '').strip()

        # 認証チェック
        username = next(
            (name for name, registered_password in ADMIN_CREDENTIALS.items()
             if registered_password == password),
            None
        )
        if username:
            session['logged_in'] = True
            session['username'] = username
            # ログイン成功後は指定されたページにリダイレクト
            return redirect(next_url)
        return render_template('login.html', error=True, message="パスワードが正しくありません。", next=next_url)

    # ログイン済みの場合は指定されたページにリダイレクト
    if session.get('logged_in'):
        return redirect(next_url)

    return render_template('login.html', next=next_url)

# ログアウト
@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('index'))

@app.route('/shelter_register', methods=['GET'])
@login_required
def shelter_register():
    mode = request.args.get('mode')
    if mode == 'new':
        return redirect(url_for('shelter_register_new'))
    if mode == 'update':
        return redirect(url_for('shelter_register_update'))
    if mode == 'delete':
        return redirect(url_for('shelter_register_delete'))
    return render_template('shelter_register_select.html', shelters=shelters)


@app.route('/shelter_register/new', methods=['GET', 'POST'])
@login_required
def shelter_register_new():
    if request.method == 'POST':
        form_data = {
            'name': request.form.get('name', '').strip(),
            'postal_code': request.form.get('postal_code', '').strip(),
            'address': request.form.get('address', '').strip(),
            'capacity': request.form.get('capacity', '0').strip(),
            'status': request.form.get('status', '空きあり'),
            'male_toilet_count': request.form.get('male_toilet_count', '0').strip(),
            'female_toilet_count': request.form.get('female_toilet_count', '0').strip(),
            'accessible_toilet_count': request.form.get('accessible_toilet_count', '0').strip(),
            'barrier_free': request.form.get('barrier_free', '未選択'),
            'pet_allowed': request.form.get('pet_allowed') is not None,
        }

        if not form_data['name']:
            return render_template('shelter_management_form.html', mode='new', error=True, message='建物名を入力してください。', shelters=shelters, form_data=form_data)

        try:
            capacity = max(0, int(float(form_data['capacity']))) if form_data['capacity'] not in ('', None) else 0
            male = max(0, int(float(form_data['male_toilet_count']))) if form_data['male_toilet_count'] not in ('', None) else 0
            female = max(0, int(float(form_data['female_toilet_count']))) if form_data['female_toilet_count'] not in ('', None) else 0
            accessible = max(0, int(float(form_data['accessible_toilet_count']))) if form_data['accessible_toilet_count'] not in ('', None) else 0
        except ValueError:
            return render_template('shelter_management_form.html', mode='new', error=True, message='収容人数・トイレ数は数値で入力してください。', shelters=shelters, form_data=form_data)

        image_url = ''
        uploaded_file = request.files.get('image')
        if uploaded_file and uploaded_file.filename:
            try:
                image_url = save_uploaded_image(uploaded_file)
            except ValueError as exc:
                return render_template('shelter_management_form.html', mode='new', error=True, message=str(exc), shelters=shelters, form_data=form_data)

        next_id = max((shelter.get('id', 0) for shelter in shelters), default=0) + 1
        shelter = {
            'id': next_id,
            'name': form_data['name'],
            'postal_code': form_data['postal_code'],
            'address': form_data['address'],
            'capacity': capacity,
            'status': form_data['status'],
            'male_toilet_count': male,
            'female_toilet_count': female,
            'accessible_toilet_count': accessible,
            'barrier_free': form_data['barrier_free'],
            'pet_allowed': form_data['pet_allowed'],
            'image': image_url,
        }
        shelters.append(shelter)
        save_shelters()
        return render_template('shelter_management_form.html', mode='new', success=True, message='避難所を登録しました。', shelters=shelters, form_data=shelter)

    return render_template('shelter_management_form.html', mode='new', shelters=shelters)


@app.route('/shelter_register/update', methods=['GET', 'POST'])
@login_required
def shelter_register_update():
    selected_id = request.args.get('shelter_id') or request.form.get('shelter_id')
    selected_shelter = get_shelter_by_id(selected_id) if selected_id else None

    if request.method == 'POST':
        shelter_id = request.form.get('shelter_id')
        shelter = get_shelter_by_id(shelter_id)
        if not shelter_id or shelter is None:
            return render_template('shelter_management_form.html', mode='update', error=True, message='更新対象を選択してください。', shelters=shelters, selected_shelter=None)

        form_data = {
            'name': request.form.get('name', '').strip(),
            'postal_code': request.form.get('postal_code', '').strip(),
            'address': request.form.get('address', '').strip(),
            'capacity': request.form.get('capacity', '0').strip(),
            'status': request.form.get('status', '空きあり'),
            'male_toilet_count': request.form.get('male_toilet_count', '0').strip(),
            'female_toilet_count': request.form.get('female_toilet_count', '0').strip(),
            'accessible_toilet_count': request.form.get('accessible_toilet_count', '0').strip(),
            'barrier_free': request.form.get('barrier_free', '未選択'),
            'pet_allowed': request.form.get('pet_allowed') is not None,
        }

        if not form_data['name']:
            return render_template('shelter_management_form.html', mode='update', error=True, message='建物名を入力してください。', shelters=shelters, selected_shelter=shelter, form_data=form_data)

        try:
            capacity = max(0, int(float(form_data['capacity']))) if form_data['capacity'] not in ('', None) else 0
            male = max(0, int(float(form_data['male_toilet_count']))) if form_data['male_toilet_count'] not in ('', None) else 0
            female = max(0, int(float(form_data['female_toilet_count']))) if form_data['female_toilet_count'] not in ('', None) else 0
            accessible = max(0, int(float(form_data['accessible_toilet_count']))) if form_data['accessible_toilet_count'] not in ('', None) else 0
        except ValueError:
            return render_template('shelter_management_form.html', mode='update', error=True, message='収容人数・トイレ数は数値で入力してください。', shelters=shelters, selected_shelter=shelter, form_data=form_data)

        existing_image = shelter.get('image', '')
        uploaded_file = request.files.get('image')
        if uploaded_file and uploaded_file.filename:
            try:
                new_image = save_uploaded_image(uploaded_file)
                if existing_image and is_local_upload_path(existing_image):
                    remove_uploaded_image(existing_image)
                shelter['image'] = new_image
            except ValueError as exc:
                return render_template('shelter_management_form.html', mode='update', error=True, message=str(exc), shelters=shelters, selected_shelter=shelter, form_data=form_data)

        shelter.update({
            'name': form_data['name'],
            'postal_code': form_data['postal_code'],
            'address': form_data['address'],
            'capacity': capacity,
            'status': form_data['status'],
            'male_toilet_count': male,
            'female_toilet_count': female,
            'accessible_toilet_count': accessible,
            'barrier_free': form_data['barrier_free'],
            'pet_allowed': form_data['pet_allowed'],
        })
        save_shelters()
        return render_template('shelter_management_form.html', mode='update', success=True, message='避難所情報を更新しました。', shelters=shelters, selected_shelter=shelter, form_data=shelter)

    if selected_id and selected_shelter is None:
        return render_template('shelter_management_form.html', mode='update', error=True, message='対象の避難所が見つかりませんでした。', shelters=shelters, selected_shelter=None)

    return render_template('shelter_management_form.html', mode='update', shelters=shelters, selected_shelter=selected_shelter, form_data=selected_shelter or {})


@app.route('/shelter_register/delete', methods=['GET', 'POST'])
@login_required
def shelter_register_delete():
    if request.method == 'POST':
        shelter_id = request.form.get('shelter_id')
        if not shelter_id:
            return render_template('shelter_management_form.html', mode='delete', error=True, message='削除する避難所を選択してください。', shelters=shelters)

        shelter = get_shelter_by_id(shelter_id)
        if shelter is None:
            return render_template('shelter_management_form.html', mode='delete', error=True, message='対象の避難所が見つかりません。', shelters=shelters)

        if request.form.get('confirm') != 'true':
            return render_template('shelter_management_form.html', mode='delete', error=True, message='削除前に確認を行ってください。', shelters=shelters, selected_shelter=shelter)

        if shelter.get('image') and is_local_upload_path(shelter.get('image')):
            remove_uploaded_image(shelter.get('image'))
        shelters[:] = [item for item in shelters if item.get('id') != shelter.get('id')]
        save_shelters()
        return render_template('shelter_management_form.html', mode='delete', success=True, message='避難所を削除しました。', shelters=shelters)

    return render_template('shelter_management_form.html', mode='delete', shelters=shelters)

# 避難所検索ページ
@app.route('/shelter_search')
def shelter_search():
    return render_template('shelter_search.html')

# 全施設一覧ページ
@app.route('/all_shelters')
def all_shelters():
    return render_template('search_results.html', results=shelters)


# 指示・発信ボード：Excel から常に最新のデータを読み込んで表示する
@app.route('/board')
@login_required
def board():
    board_data = load_board_data()
    return render_template(
        'board.html',
        board_data=board_data,
        damage_map_points=build_damage_map_points(board_data.get('被害状況一覧', [])),
        map_center=MAP_CENTER,
    )

# 検索結果ページ：templates/search_results.html を返す
@app.route('/search_results')
def search_results():
    pet_allowed = request.args.get('pet_allowed') or request.args.get('pet-allowed')
    barrier_free = request.args.get('barrier_free') or request.args.get('barrier-free')
    filters = {
        'name': request.args.get('name', ''),
        'postal_code': request.args.get('postal_code', ''),
        'address': request.args.get('address', ''),
        'capacity': request.args.get('capacity', ''),
        'pet_allowed': pet_allowed,
        'barrier_free': barrier_free,
    }
    results = search_shelters(filters)
    return render_template('search_results.html', results=results, filters=filters)

# JSON API：/shelters?district=地区名
@app.route('/shelters', methods=['GET'])
def get_shelters():
    results = filter_shelters(request.args.get('district'))

    if not results:
        # 見つからなければエラー JSON を返す
        return jsonify({'error': 'No shelters found'}), 404

    # 見つかったらリストを JSON で返す
    return jsonify(results)


# 気象警報・注意報API
@app.route('/api/weather_warnings')
def api_weather_warnings():
    """気象警報・注意報をJSON形式で返すAPI"""
    return jsonify(get_weather_warnings(force_refresh=request.args.get('refresh') == '1'))

if __name__ == '__main__':
    app.run(debug=True, port=5000)
